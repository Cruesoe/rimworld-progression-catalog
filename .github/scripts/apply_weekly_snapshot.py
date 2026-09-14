#!/usr/bin/env python3
"""Refresh the Progression catalog from the public Steam collection APIs."""
from __future__ import annotations

import csv
import json
import re
import urllib.parse
import urllib.request
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
HTML = ROOT / "index.html"
README = ROOT / "README.md"
CSV_PATH = ROOT / "backup/progression_packageids.csv"
TXT_PATH = ROOT / "backup/progression_collections_categorized.txt"
MODPACK_PATH = ROOT / "backup/progression_modpack_mods.txt"

COLLECTIONS = {
    "core": ("3521297585", "1/3 Core"),
    "content": ("3521319712", "2/3 Content"),
    "cosmetics": ("3637541646", "3/3 Cosmetics"),
}
COLLECTION_API = "https://api.steampowered.com/ISteamRemoteStorage/GetCollectionDetails/v1/"
DETAILS_API = "https://api.steampowered.com/ISteamRemoteStorage/GetPublishedFileDetails/v1/"


def post_json(url: str, fields: dict[str, str]) -> dict:
    request = urllib.request.Request(
        url,
        data=urllib.parse.urlencode(fields).encode("ascii"),
        headers={"User-Agent": "progression-catalog-weekly-refresh/1.0"},
    )
    with urllib.request.urlopen(request, timeout=60) as response:
        return json.load(response)


def collection_ids(collection_id: str) -> set[str]:
    payload = post_json(COLLECTION_API, {"collectioncount": "1", "publishedfileids[0]": collection_id})
    details = payload.get("response", {}).get("collectiondetails", [])
    if not details or int(details[0].get("result", 0)) != 1:
        raise RuntimeError(f"Steam did not return collection {collection_id}")
    return {str(child["publishedfileid"]) for child in details[0].get("children", [])}


def published_details(ids: set[str]) -> dict[str, dict]:
    result: dict[str, dict] = {}
    ordered = sorted(ids, key=int)
    for start in range(0, len(ordered), 100):
        batch = ordered[start : start + 100]
        fields = {"itemcount": str(len(batch))}
        fields.update({f"publishedfileids[{i}]": item_id for i, item_id in enumerate(batch)})
        payload = post_json(DETAILS_API, fields)
        for item in payload.get("response", {}).get("publishedfiledetails", []):
            result[str(item["publishedfileid"])] = item
    missing = ids - result.keys()
    if missing:
        raise RuntimeError(f"Steam omitted {len(missing)} requested workshop items")
    return result


def load_data(html: str) -> dict:
    match = re.search(r"const DATA = (\{.*?\});\s*\n", html, re.S)
    if not match:
        raise RuntimeError("DATA block not found")
    return json.loads(match.group(1))


def infer_category(title: str, categories: list[str]) -> str:
    text = title.casefold()
    rules = [
        ("Frameworks & libraries", ("framework", "library", " lib")),
        ("Performance & tools", ("performance", "optimizer", "loading", "fps ", "tool")),
        ("UI & information", (" ui", "tab", "menu", "tooltip", "overlay", "search")),
        ("Textures & retextures", ("texture", "retexture")),
        ("Hair, bodies & animations", ("hair", "body", "animation")),
        ("Apparel, styles & cosmetics", ("apparel", "clothing", "style", "cosmetic")),
        ("Audio & music", ("audio", "music", "sound")),
        ("Patches & compatibility", ("patch", "compat", "fix")),
        ("Vanilla Expanded", ("vanilla expanded", " vfe", " vre", "vpe ")),
        ("Alpha series", ("alpha ",)),
        ("Animals", ("animal", " dog", " cat", " livestock")),
        ("Combat, weapons & raids", ("weapon", "combat", "raid", "turret", "armor")),
        ("Building, production & storage", ("building", "production", "storage", "furniture")),
        ("Medical, needs & hygiene", ("medical", "health", "hygiene", "hospital")),
        ("Genes, xenotypes & races", ("gene", "xenotype", "race")),
        ("Ideology & social", ("ideology", "social", "ritual")),
        ("World, biomes, factions & travel", ("world", "biome", "faction", "travel")),
        ("Research, storytellers & progression", ("research", "storyteller", "progression")),
        ("Vehicles, gravships & mechs", ("vehicle", "gravship", "mech")),
    ]
    for category, needles in rules:
        if category in categories and any(needle in text for needle in needles):
            return category
    return "Gameplay tweaks" if "Gameplay tweaks" in categories else categories[-1]


def refresh_items(data: dict) -> tuple[dict, dict]:
    memberships = {key: collection_ids(collection_id) for key, (collection_id, _) in COLLECTIONS.items()}
    collection_page_ids = {collection_id for collection_id, _ in COLLECTIONS.values()}
    for item_ids in memberships.values():
        item_ids.difference_update(collection_page_ids)
    current_ids = set().union(*memberships.values())
    details = published_details(current_ids)
    previous = {str(item["id"]): item for item in data["items"]}
    changes = {"added": [], "removed": [], "renamed": [], "newly_unavailable": []}
    for item_id, item in previous.items():
        if item_id not in current_ids:
            changes["removed"].append({"t": item["t"], "id": item_id})

    refreshed = []
    for item_id in current_ids:
        old = previous.get(item_id)
        detail = details[item_id]
        available = int(detail.get("result", 0)) == 1
        title = str(detail.get("title") or (old or {}).get("t") or f"Workshop item {item_id}") if available else f"(unavailable {item_id})"
        if old is None:
            changes["added"].append({"t": title, "id": item_id})
        elif available and not old.get("gone") and old["t"] != title:
            changes["renamed"].append({"from": old["t"], "to": title, "id": item_id})
        elif not available and not old.get("gone"):
            changes["newly_unavailable"].append({"t": old["t"], "id": item_id})
        item = {
            "t": title,
            "c": (old or {}).get("c") or infer_category(title, data["cats"]),
            "core": item_id in memberships["core"],
            "content": item_id in memberships["content"],
            "cosmetics": item_id in memberships["cosmetics"],
            "id": item_id,
            "u": f"https://steamcommunity.com/sharedfiles/filedetails/?id={item_id}",
            "p": (old or {}).get("p", ""),
        }
        if not available:
            item["gone"] = True
        refreshed.append(item)
    order = {category: index for index, category in enumerate(data["cats"])}
    refreshed.sort(key=lambda item: (order.get(item["c"], len(order)), item["t"].casefold(), int(item["id"])))
    data["items"] = refreshed
    return data, changes


def change_summary(changes: dict) -> str:
    return f"{len(changes['added'])} added · {len(changes['removed'])} removed · {len(changes['renamed'])} renamed · {len(changes['newly_unavailable'])} newly unavailable"


def update_html(html: str, data: dict, changes: dict, iso_date: str, display_date: str) -> str:
    html = re.sub(r'<div class="updated">.*?</div>', f'<div class="updated">Steam snapshot: <time datetime="{iso_date}">{display_date}</time></div>', html, count=1, flags=re.S)
    html = re.sub(r'(<details class="changes">\s*<summary>).*?(</summary>)', r'\g<1>Latest snapshot changes: ' + change_summary(changes) + r'\g<2>', html, count=1, flags=re.S)
    html = re.sub(r"const DATA = \{.*?\};\s*\n", "const DATA = " + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n", html, count=1, flags=re.S)
    total = len(data["items"])
    known = sum(bool(item.get("p")) for item in data["items"])
    footer_match = re.search(r"<footer>.*?</footer>", html, re.S)
    if footer_match:
        footer = footer_match.group(0)
        footer = re.sub(r"Steam snapshot: .*?\.", f"Steam snapshot: {display_date}.", footer)
        footer = re.sub(r"(packageId from RimSort steamDB: \d+ of )\d+\.", rf"\g<1>{total}.", footer)
        footer = re.sub(r"Known packageIds \(RimSort \+ previous overlay\): \d+ of \d+\.", f"Known packageIds (RimSort + previous overlay): {known} of {total}.", footer)
        html = html[:footer_match.start()] + footer + html[footer_match.end():]
    return html


def format_change_list(changes: dict) -> str:
    sections = []
    labels = (("added", "Added"), ("removed", "Removed"), ("renamed", "Renamed"), ("newly_unavailable", "Newly unavailable"))
    for key, label in labels:
        entries = changes[key]
        sections.append(f"### {label} ({len(entries)})\n")
        if not entries:
            sections.append("- None\n")
        elif key == "renamed":
            sections.extend(f"- {entry['from']} → {entry['to']} — {entry['id']}\n" for entry in entries)
        else:
            sections.extend(f"- {entry['t']} — {entry['id']}\n" for entry in entries)
        sections.append("\n")
    return "".join(sections).rstrip()


def update_readme(text: str, data: dict, changes: dict, display_date: str) -> str:
    counts = Counter()
    for item in data["items"]:
        for key in COLLECTIONS:
            if item[key]:
                counts[key] += 1
    total = len(data["items"])
    known = sum(bool(item.get("p")) for item in data["items"])
    text = re.sub(r"\*\*Steam snapshot date:.*?\*\*", f"**Steam snapshot date: {display_date}**", text, count=1)
    text = re.sub(r"(\*\*packageId from RimSort steamDB: \d+ of )\d+(\*\*)", rf"\g<1>{total}\g<2>", text, count=1)
    text = re.sub(r"\*\*Known packageIds .*?\*\*", f"**Known packageIds (RimSort + previous overlay): {known} of {total}**", text, count=1)
    text = re.sub(r"\*\*Unique mods:.*?\*\*[^\n]*", f"**Unique mods: {total}** — Core {counts['core']} · Content {counts['content']} · Cosmetics {counts['cosmetics']} (3 collection pages omitted from the item list)", text, count=1)
    replacement = "## This snapshot vs previous catalog\n\n" + format_change_list(changes) + "\n\n"
    return re.sub(r"## This snapshot vs previous catalog\n.*?(?=## Open the catalog)", replacement, text, count=1, flags=re.S)


def write_backups(data: dict, display_date: str) -> None:
    with CSV_PATH.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["title", "workshop_id", "packageId", "collection_core", "collection_content", "collection_cosmetics", "category", "steam_url"])
        for item in sorted(data["items"], key=lambda value: (value["t"].casefold(), value["id"])):
            writer.writerow([item["t"], item["id"], item.get("p", ""), int(item["core"]), int(item["content"]), int(item["cosmetics"]), item["c"], item["u"]])
    lines = ["The Progression — categorized catalog", f"Steam snapshot: {display_date}", ""]
    for category in data["cats"]:
        items = [item for item in data["items"] if item["c"] == category]
        if not items:
            continue
        lines += ["=" * 72, f"{category}  ({len(items)})", "=" * 72]
        for item in items:
            tags = [label for key, (_, label) in COLLECTIONS.items() if item[key]]
            unavailable = "  [unavailable]" if item.get("gone") else ""
            lines += [f"  - {item['t']}{unavailable}", f"    [{' · '.join(tags)}]  {item['u']}"]
        lines.append("")
    TXT_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")
    core = [item for item in data["items"] if item["core"]]
    lines = ["The Progression Modpack [1.6]", "https://steamcommunity.com/workshop/filedetails/?id=3521297585", f"Total items: {len(core)}", ""]
    for number, item in enumerate(core, 1):
        lines += [f"   {number}. {item['t']}", f"     {item['u']}"]
    MODPACK_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl missing; skipped xlsx")
        return
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["A1"] = "The Progression collections — categorized workshop list"
    summary["A3"] = f"Steam snapshot: {display_date}"
    mods = workbook.create_sheet("All mods")
    mods.append(["Category", "Mod title", "Collections", "Core 1/3", "Content 2/3", "Cosmetics 3/3", "Workshop ID", "Steam URL", "packageId", "Unavailable"])
    head = Font(name="Arial", bold=True, size=11, color="E8F5EF")
    fill = PatternFill("solid", fgColor="1B3A2E")
    for column in range(1, 11):
        mods.cell(1, column).font = head
        mods.cell(1, column).fill = fill
    for item in data["items"]:
        tags = [label for key, (_, label) in COLLECTIONS.items() if item[key]]
        mods.append([item["c"], item["t"], " · ".join(tags), "Yes" if item["core"] else None, "Yes" if item["content"] else None, "Yes" if item["cosmetics"] else None, item["id"], item["u"], item.get("p") or None, "Yes" if item.get("gone") else None])
    for column, width in enumerate([36, 55, 28, 12, 14, 16, 16, 62, 36, 14], 1):
        mods.column_dimensions[get_column_letter(column)].width = width
    workbook.save(ROOT / "backup/progression_collections_categorized.xlsx")


def main() -> None:
    now = datetime.now(timezone.utc)
    iso_date = now.strftime("%Y-%m-%d")
    display_date = f"{now.day} {now.strftime('%B %Y')}"
    html = HTML.read_text(encoding="utf-8")
    data, changes = refresh_items(load_data(html))
    HTML.write_text(update_html(html, data, changes, iso_date, display_date), encoding="utf-8")
    README.write_text(update_readme(README.read_text(encoding="utf-8"), data, changes, display_date), encoding="utf-8")
    write_backups(data, display_date)
    print(f"refreshed {len(data['items'])} mods: {change_summary(changes)}")


if __name__ == "__main__":
    main()
